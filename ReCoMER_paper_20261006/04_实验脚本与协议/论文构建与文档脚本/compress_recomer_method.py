from docx import Document
from pathlib import Path

src = Path(r"C:\Users\肖田泽宇宙最强1234\Desktop\论文\ReCoMER_Method_中英双语_20261004_最终版.docx")
dst = Path(r"C:\Users\肖田泽宇宙最强1234\Desktop\论文\ReCoMER_Method_中英双语_20261004_2700词版.docx")

doc = Document(str(src))

def replace_in_paragraph(index: int, replacements):
    p = doc.paragraphs[index]
    text = p.text
    for old, new in replacements:
        if old not in text:
            raise ValueError(f"Paragraph {index}: missing text: {old!r}")
        text = text.replace(old, new)
    # These are prose-only paragraphs; equation paragraphs are not touched.
    if not p.runs:
        p.add_run(text)
    else:
        p.runs[0].text = text
        for run in p.runs[1:]:
            run.text = ""

# 3.2--3.4: remove repetition while retaining every implementation constraint,
# formula interpretation, and the closed-loop rationale.
replace_in_paragraph(7, [
    (" This masking is important because the model is intended to work with incomplete multimodal inputs rather than treating zero padding as meaningful evidence.", ""),
])
replace_in_paragraph(9, [
    (" These heads are not the final classifier.", ""),
    (" This order prevents the router from assigning a weight to a stale representation that was computed before historical evidence was considered.", ""),
])
replace_in_paragraph(11, [
    (" The final predictor can also be evaluated with a controlled modality subset, which is required for the Shapley teacher described in Section 3.3.", ""),
])
replace_in_paragraph(13, [
    (" A masked mean of the available projections produces a shared context vector.", " A masked mean produces the shared context."),
])
replace_in_paragraph(14, [
    (" The seven evaluations also make the target sensitive to complementary evidence.", ""),
])
replace_in_paragraph(16, [
    (" If all modalities have similar marginal utility, the standardized target is close to uniform; if one modality supplies distinctive evidence, its target becomes positive relative to the others.", ""),
])
replace_in_paragraph(18, [
    (" This range prevents a single noisy modality from becoming arbitrarily dominant and also prevents a weak but present modality from being assigned a negative contribution.", ""),
])
replace_in_paragraph(20, [
    (" Equal contribution scores recover uniform weighting, which provides a direct and interpretable control condition for ablation.", " Equal scores recover uniform weighting for ablation."),
])
replace_in_paragraph(22, [
    (" Three text semantic-similarity positions are reserved in the feature layout but are set to zero in the active base configuration.", " The three reserved text-similarity positions are zero in the base configuration."),
])
replace_in_paragraph(24, [
    (" This construction preserves the original history candidates instead of replacing them with a single scalar gate.", ""),
])
replace_in_paragraph(26, [
    (" This distinction makes the abstention mechanism easier to interpret: the null candidate represents a decision not to use history, whereas a real slot represents a concrete previous utterance.", " The null represents abstention, whereas a real slot represents a concrete previous utterance."),
])

# 3.5--3.6: compact repeated motivation while retaining the independent expert,
# class-wise evidence, candidate safety, and recoverable-baseline guarantees.
replace_in_paragraph(30, [
    (" The TAV branch supplies the reference distribution because it covers all available modalities.", " TAV supplies the all-modality reference."),
])
replace_in_paragraph(31, [
    (" Audio and vision are projected independently to 256 dimensions and layer-normalized.", " Audio and vision are independently projected to 256 dimensions and layer-normalized."),
])
replace_in_paragraph(32, [
    ("The split between TAV and AV is also a constraint: the AV branch must remain text-free if it is to provide independent nonverbal evidence.", "The AV branch remains text-free to preserve independent nonverbal evidence."),
])
replace_in_paragraph(33, [
    (" The gate does not choose a single branch. Instead, it controls the strength of the complete class-wise evidence vector, allowing several classes to be increased or decreased at once.", " The gate scales the complete class-wise evidence vector, so several classes can be increased or decreased together."),
])
replace_in_paragraph(38, [
    (" The first two are the peer predictions that define the baseline; the last two provide anchors for measuring how much the complete predictions changed relative to their internal references.", " The first two define the peer baseline; the last two anchor changes from each internal reference."),
])
replace_in_paragraph(40, [
    (" The separation between P and C is intentional.", " P and C remain separate by design."),
])
replace_in_paragraph(42, [
    (" Feature means and standard deviations are fitted only on the fusion fitting set, and standard deviations are clamped to ", " Feature statistics are fitted only on the fusion set and clamped to "),
])
replace_in_paragraph(44, [
    (" This rule prevents the corrector from transferring mass to an arbitrary low-probability class based on a small numerical fluctuation.", ""),
])
replace_in_paragraph(46, [
    (" The preservation term discourages unnecessary changes to correct decisions but does not assume that every baseline decision is correct.", " The preservation term limits unnecessary changes without assuming every baseline decision is correct."),
])

doc.save(str(dst))
print(dst)
