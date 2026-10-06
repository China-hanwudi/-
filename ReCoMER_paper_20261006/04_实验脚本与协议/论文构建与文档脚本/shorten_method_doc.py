from pathlib import Path
import re
from docx import Document


SRC = Path(r"C:\Users\肖田泽宇宙最强1234\Desktop\论文\ReCoMER_Method_中英双语_20261004_逻辑修订版.docx")
OUT = SRC.with_name("ReCoMER_Method_中英双语_20261004_最终版.docx")


def split_sentences(text):
    return re.split(r"(?<=[.!?])\s+", text.strip())


def set_text(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def main():
    doc = Document(SRC)
    # These paragraphs do not carry the formula-introduction suffixes. Removing
    # one redundant closing sentence keeps the English 3.2--3.6 block within
    # the requested 2,500--3,000 words without deleting implementation facts.
    for idx in (11, 13, 20, 22, 28, 30, 31, 32, 35, 36, 40, 48):
        text = doc.paragraphs[idx].text
        parts = split_sentences(text)
        if len(parts) > 3:
            set_text(doc.paragraphs[idx], " ".join(parts[:-1]))
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
